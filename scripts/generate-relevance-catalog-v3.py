#!/usr/bin/env python3
"""Validate MetaMed relevance v3 and generate its deterministic SQL migration.

The workbook is read-only. The generator validates source rows, cached formula
results and business invariants before writing SQL. It never connects to a
database and never edits the source workbook.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import unicodedata
import uuid
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Iterable

import openpyxl


PUBLICATION_KEY = "metamed-relevancia-global-2026-v3"
UUID_NAMESPACE = uuid.UUID("7377453b-2d46-4b95-9af2-7c6f171eb53f")
AUDITED_WORKBOOK_SHA256 = "3671a6c8ce245c5e89e744671adc518b48cef4151b195352505dc2819da3fd7e"
AUDITED_DOCUMENT_SHA256 = "7710cf33765ec6d5145ab15cecff2d21f3c803af015c7949172835f54188f668"
EXPECTED_TOPIC_COUNT = 492
EXPECTED_ITEM_COUNTS = {"Medway": 172, "MedGrupo": 90, "MedCof": 658}
EXPECTED_LINK_COUNTS = {"Medway": 324, "MedGrupo": 226, "MedCof": 685}
EXPECTED_HEADERS = {
    "temas_canonicos": [
        "tema_id", "tema_canonico", "grande_area", "especialidade",
        "relevancia_global", "sinal_medcof", "confianca", "justificativa",
        "sinonimos", "ano_referencia",
    ],
    "mapping": [
        "id_item", "cursinho", "area_no_cursinho", "assunto_no_cursinho",
        "tema_id", "tema_canonico", "peso", "nota_tema", "peso_x_nota",
        "media_ponderada", "nota_maxima", "relevancia_efetiva",
    ],
}
SHEET_PROVIDER = {
    "depara_medway": "Medway",
    "depara_medgrupo": "MedGrupo",
    "depara_medcof": "MedCof",
}
AREA_NORMALIZATION = {
    "Clínica Médica": "Clínica Médica",
    "Cirurgia Geral": "Cirurgia",
    "Pediatria": "Pediatria",
    "Ginecologia e Obstetrícia": "Ginecologia e Obstetrícia",
    "Medicina Preventiva": "Preventiva",
}
AREA_BY_PREFIX = {
    "CM": "Clínica Médica",
    "CIR": "Cirurgia Geral",
    "PED": "Pediatria",
    "GO": "Ginecologia e Obstetrícia",
    "MP": "Medicina Preventiva",
}
RELEASE_META = {
    "Medway": {
        "provider_code": "medway",
        "release_name": "Extensivo R1 2026",
        "source_document": "Conteúdo programático, PDF institucional",
        "source_edition": "Extensivo R1 2026",
        "source_volume": "185 aulas",
        "source_extracted_on": None,
    },
    "MedGrupo": {
        "provider_code": "medgrupo",
        "release_name": "Medcurso 2020",
        "source_document": "Cronograma Medcurso",
        "source_edition": "2020",
        "source_volume": "46 semanas, 2 temas por semana",
        "source_extracted_on": None,
    },
    "MedCof": {
        "provider_code": "medcof",
        "release_name": "Extensivo 2025 — R1 Acesso Direto",
        "source_document": "Cronograma da plataforma aulas.medcof.com.br",
        "source_edition": "Extensivo 2025 — R1 Acesso Direto",
        "source_volume": "640 aulas",
        "source_extracted_on": date(2026, 9, 6),
    },
}


class ValidationError(RuntimeError):
    pass


@dataclass(frozen=True)
class Topic:
    topic_id: str
    canonical_name: str
    source_major_area: str
    major_area: str
    specialty: str
    score: int
    editorial_signal: str | None
    confidence: str
    rationale: str | None
    aliases: tuple[str, ...]
    reference_year: int


@dataclass(frozen=True)
class Link:
    topic_id: str
    weight: Decimal


@dataclass(frozen=True)
class CatalogItem:
    provider: str
    external_id: str
    catalog_area: str
    title: str
    source_order: int
    links: tuple[Link, ...]
    effective_score: Decimal
    major_area: str
    specialty: str


@dataclass(frozen=True)
class Dataset:
    workbook_sha256: str
    document_sha256: str
    topics: dict[str, Topic]
    items: dict[str, tuple[CatalogItem, ...]]
    formula_count: int
    alias_count: int


def fail(condition: bool, message: str) -> None:
    if not condition:
        raise ValidationError(message)


def text_value(value: Any, field: str) -> str:
    fail(value is not None and str(value).strip() != "", f"Missing {field}")
    return str(value).strip()


def normalize_alias(value: str) -> str:
    decomposed = unicodedata.normalize("NFD", value)
    without_marks = "".join(ch for ch in decomposed if unicodedata.category(ch) != "Mn")
    return re.sub(r"\s+", " ", without_marks).strip().casefold()


def decimal_value(value: Any, field: str) -> Decimal:
    try:
        return Decimal(str(value))
    except Exception as exc:  # pragma: no cover - defensive input error detail
        raise ValidationError(f"Invalid decimal in {field}: {value!r}") from exc


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def uuid_for(kind: str, *parts: str) -> uuid.UUID:
    return uuid.uuid5(UUID_NAMESPACE, ":".join((kind, *parts)))


def check_headers(sheet: Any, expected: list[str]) -> None:
    actual = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    fail(actual == expected, f"Unexpected headers in {sheet.title}: {actual!r}")


def load_topics(cached_sheet: Any) -> dict[str, Topic]:
    check_headers(cached_sheet, EXPECTED_HEADERS["temas_canonicos"])
    topics: dict[str, Topic] = {}
    prefix_counts: Counter[tuple[str, str]] = Counter()
    for row_number, row in enumerate(cached_sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not any(value is not None for value in row):
            continue
        topic_id = text_value(row[0], f"temas_canonicos!A{row_number}")
        fail(topic_id not in topics, f"Duplicate topic_id {topic_id}")
        match = re.fullmatch(r"(CM|CIR|PED|GO|MP)-[A-Z0-9]+-\d{3}", topic_id)
        fail(match is not None, f"Invalid topic_id format: {topic_id}")
        source_area = text_value(row[2], f"temas_canonicos!C{row_number}")
        fail(source_area in AREA_NORMALIZATION, f"Unknown canonical major area: {source_area}")
        fail(source_area == AREA_BY_PREFIX[match.group(1)], f"Area/prefix mismatch for {topic_id}: {source_area}")
        score_decimal = decimal_value(row[4], f"temas_canonicos!E{row_number}")
        fail(score_decimal == score_decimal.to_integral_value(), f"Non-integer relevance for {topic_id}")
        score = int(score_decimal)
        fail(1 <= score <= 10, f"Relevance outside 1..10 for {topic_id}")
        reference_year = int(row[9])
        fail(reference_year == 2026, f"Unexpected reference year for {topic_id}: {reference_year}")
        aliases: list[str] = []
        seen_aliases: set[str] = set()
        for alias in str(row[8] or "").split(";"):
            alias = alias.strip()
            if not alias:
                continue
            key = normalize_alias(alias)
            fail(key not in seen_aliases, f"Duplicate normalized alias for {topic_id}: {alias}")
            seen_aliases.add(key)
            aliases.append(alias)
        topics[topic_id] = Topic(
            topic_id=topic_id,
            canonical_name=text_value(row[1], f"temas_canonicos!B{row_number}"),
            source_major_area=source_area,
            major_area=AREA_NORMALIZATION[source_area],
            specialty=text_value(row[3], f"temas_canonicos!D{row_number}"),
            score=score,
            editorial_signal=str(row[5]).strip() if row[5] not in (None, "") else None,
            confidence=text_value(row[6], f"temas_canonicos!G{row_number}"),
            rationale=str(row[7]).strip() if row[7] not in (None, "") else None,
            aliases=tuple(aliases),
            reference_year=reference_year,
        )
        prefix_counts[(match.group(1), source_area)] += 1
    fail(len(topics) == EXPECTED_TOPIC_COUNT, f"Expected 492 topics, found {len(topics)}")
    fail(len(prefix_counts) == 5, f"Canonical area correction is incomplete: {prefix_counts}")
    return topics


def close_decimal(actual: Any, expected: Decimal, field: str, tolerance: Decimal = Decimal("0.000001")) -> None:
    value = decimal_value(actual, field)
    fail(abs(value - expected) <= tolerance, f"Cached result mismatch in {field}: {value} != {expected}")


def load_items(
    formula_sheet: Any,
    cached_sheet: Any,
    provider: str,
    topics: dict[str, Topic],
) -> tuple[tuple[CatalogItem, ...], int]:
    check_headers(formula_sheet, EXPECTED_HEADERS["mapping"])
    check_headers(cached_sheet, EXPECTED_HEADERS["mapping"])
    groups: dict[str, list[tuple[int, tuple[Any, ...], tuple[Any, ...]]]] = defaultdict(list)
    formula_count = 0
    formula_rows = formula_sheet.iter_rows(min_row=2)
    cached_rows = cached_sheet.iter_rows(min_row=2, values_only=True)
    for row_number, (formula_row, cached_row) in enumerate(zip(formula_rows, cached_rows), start=2):
        formula_values = tuple(cell.value for cell in formula_row)
        if not any(value is not None for value in formula_values):
            continue
        fail(len(formula_values) >= 12 and len(cached_row) >= 12, f"Short row in {formula_sheet.title}!{row_number}")
        for index in (5, 7, 8, 9, 10, 11):
            fail(formula_row[index].data_type == "f", f"Expected formula in {formula_sheet.title}!{formula_row[index].coordinate}")
            fail(cached_row[index] is not None, f"Missing cached formula value in {cached_sheet.title}!{formula_row[index].coordinate}")
            formula_count += 1
        item_id = text_value(formula_values[0], f"{formula_sheet.title}!A{row_number}")
        groups[item_id].append((row_number, formula_values, cached_row))

    items: list[CatalogItem] = []
    pair_seen: set[tuple[str, str]] = set()
    for source_order, (item_id, rows) in enumerate(groups.items(), start=1):
        metadata = {
            (
                text_value(row[1][1], f"{formula_sheet.title}!B{row[0]}"),
                text_value(row[1][2], f"{formula_sheet.title}!C{row[0]}"),
                text_value(row[1][3], f"{formula_sheet.title}!D{row[0]}"),
            )
            for row in rows
        }
        fail(len(metadata) == 1, f"Inconsistent item attributes for {provider} {item_id}: {metadata}")
        item_provider, catalog_area, title = next(iter(metadata))
        fail(item_provider == provider, f"Provider mismatch for {item_id}: {item_provider}")
        links: list[Link] = []
        weighted_mean = Decimal("0")
        scores: list[int] = []
        for row_number, formula_values, cached_values in rows:
            topic_id = text_value(formula_values[4], f"{formula_sheet.title}!E{row_number}")
            fail(topic_id in topics, f"Unknown topic_id {topic_id} in {formula_sheet.title}!E{row_number}")
            fail((item_id, topic_id) not in pair_seen, f"Duplicate item/topic pair: {provider} {item_id} {topic_id}")
            pair_seen.add((item_id, topic_id))
            weight = decimal_value(formula_values[6], f"{formula_sheet.title}!G{row_number}")
            fail(weight > 0 and weight <= 1, f"Invalid weight for {provider} {item_id} {topic_id}: {weight}")
            topic = topics[topic_id]
            contribution = weight * Decimal(topic.score)
            weighted_mean += contribution
            scores.append(topic.score)
            links.append(Link(topic_id, weight))
            fail(str(cached_values[5]).strip() == topic.canonical_name, f"Cached canonical name mismatch in {cached_sheet.title}!F{row_number}")
            close_decimal(cached_values[7], Decimal(topic.score), f"{cached_sheet.title}!H{row_number}")
            close_decimal(cached_values[8], contribution, f"{cached_sheet.title}!I{row_number}")
        total_weight = sum((link.weight for link in links), Decimal("0"))
        fail(abs(total_weight - Decimal("1")) <= Decimal("0.000001"), f"Weights do not sum to 1 for {provider} {item_id}: {total_weight}")
        max_score = Decimal(max(scores))
        effective = (Decimal("0.7") * weighted_mean + Decimal("0.3") * max_score).quantize(Decimal("0.1"), rounding=ROUND_HALF_UP)
        for row_number, _formula_values, cached_values in rows:
            close_decimal(cached_values[9], weighted_mean, f"{cached_sheet.title}!J{row_number}")
            close_decimal(cached_values[10], max_score, f"{cached_sheet.title}!K{row_number}")
            close_decimal(cached_values[11], effective, f"{cached_sheet.title}!L{row_number}")
        # Experimental presentation label only. Aggregate reporting must follow
        # canonical topic links rather than this single-item convenience field.
        dominant = sorted(links, key=lambda link: (-link.weight, link.topic_id))[0]
        major_area = topics[dominant.topic_id].major_area
        specialties = {topics[link.topic_id].specialty for link in links}
        specialty = next(iter(specialties)) if len(specialties) == 1 else "Conteúdo integrado"
        items.append(CatalogItem(
            provider=provider,
            external_id=item_id,
            catalog_area=catalog_area,
            title=title,
            source_order=source_order,
            links=tuple(links),
            effective_score=effective,
            major_area=major_area,
            specialty=specialty,
        ))
    fail(len(items) == EXPECTED_ITEM_COUNTS[provider], f"Unexpected item count for {provider}: {len(items)}")
    fail(sum(len(item.links) for item in items) == EXPECTED_LINK_COUNTS[provider], f"Unexpected link count for {provider}")
    return tuple(items), formula_count


def load_dataset(workbook_path: Path, document_path: Path) -> Dataset:
    fail(workbook_path.is_file(), f"Workbook not found: {workbook_path}")
    fail(document_path.is_file(), f"Technical document not found: {document_path}")
    workbook_digest = sha256(workbook_path)
    document_digest = sha256(document_path)
    fail(workbook_digest == AUDITED_WORKBOOK_SHA256,
         f"Workbook SHA-256 changed; create a new publication key instead of reusing v3: {workbook_digest}")
    fail(document_digest == AUDITED_DOCUMENT_SHA256,
         f"Technical-document SHA-256 changed; create a new publication key instead of reusing v3: {document_digest}")
    formula_book = openpyxl.load_workbook(workbook_path, read_only=True, data_only=False)
    cached_book = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    expected_sheets = ["leia-me", "temas_canonicos", *SHEET_PROVIDER]
    fail(formula_book.sheetnames == expected_sheets, f"Unexpected workbook sheets: {formula_book.sheetnames}")
    fail(cached_book.sheetnames == expected_sheets, f"Unexpected cached workbook sheets: {cached_book.sheetnames}")
    topics = load_topics(cached_book["temas_canonicos"])
    items: dict[str, tuple[CatalogItem, ...]] = {}
    formula_count = 0
    for sheet_name, provider in SHEET_PROVIDER.items():
        provider_items, provider_formula_count = load_items(
            formula_book[sheet_name], cached_book[sheet_name], provider, topics
        )
        items[provider] = provider_items
        formula_count += provider_formula_count
    fail(formula_count == 7410, f"Expected 7,410 mapping formulas, found {formula_count}")
    all_ids = {item.external_id for provider_items in items.values() for item in provider_items}
    fail("MC-186" not in all_ids, "MC-186 must remain an unused tombstoned identifier")
    fail("MC-375" in all_ids, "MC-375 (Oncocirurgia: Melanoma) is missing")
    fail({"MC-417", "MC-574"}.issubset(all_ids), "Verified DSD item pair is missing")
    examples = {item.external_id: item.effective_score for provider_items in items.values() for item in provider_items}
    fail(examples.get("MW-045") == Decimal("7.7"), f"MW-045 expected 7.7, found {examples.get('MW-045')}")
    fail(examples.get("MG-058") == Decimal("7.3"), f"MG-058 expected 7.3, found {examples.get('MG-058')}")
    alias_count = sum(len(topic.aliases) for topic in topics.values())
    return Dataset(
        workbook_sha256=workbook_digest,
        document_sha256=document_digest,
        topics=topics,
        items=items,
        formula_count=formula_count,
        alias_count=alias_count,
    )


def sql_literal(value: Any) -> str:
    if value is None:
        return "null"
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (int, Decimal)):
        return str(value)
    if isinstance(value, date):
        return "'" + value.isoformat() + "'::date"
    return "'" + str(value).replace("'", "''") + "'"


def insert_rows(table: str, columns: Iterable[str], rows: Iterable[Iterable[Any]], chunk_size: int = 200) -> str:
    columns = list(columns)
    materialized = [list(row) for row in rows]
    statements: list[str] = []
    for offset in range(0, len(materialized), chunk_size):
        chunk = materialized[offset:offset + chunk_size]
        values = ",\n".join("  (" + ", ".join(sql_literal(value) for value in row) + ")" for row in chunk)
        statements.append(f"insert into {table} ({', '.join(columns)}) values\n{values};")
    return "\n\n".join(statements)


def render_migration(dataset: Dataset, workbook_path: Path, document_path: Path) -> str:
    publication_id = uuid_for("publication", PUBLICATION_KEY)
    releases = {
        provider: uuid_for("release", PUBLICATION_KEY, RELEASE_META[provider]["provider_code"], RELEASE_META[provider]["release_name"])
        for provider in RELEASE_META
    }
    schema = f"""-- Generated by scripts/generate-relevance-catalog-v3.py after full source validation.
-- Source workbook SHA-256: {dataset.workbook_sha256}
-- Technical document SHA-256: {dataset.document_sha256}
-- Safety: v3 is seeded as validated, never active. No existing schedule changes.

create table public.relevance_publications (
  id uuid primary key,
  version_key text not null unique,
  reference_year smallint not null check (reference_year between 2000 and 2200),
  validation_state text not null check (validation_state in ('draft','validated','rejected')),
  score_formula text not null,
  rounding_mode text not null check (rounding_mode = 'half_up_1_decimal'),
  source_filename text not null,
  source_sha256 text not null check (source_sha256 ~ '^[0-9a-f]{{64}}$'),
  technical_document_filename text not null,
  technical_document_sha256 text not null check (technical_document_sha256 ~ '^[0-9a-f]{{64}}$'),
  importer_version text not null,
  validation_report jsonb not null check (jsonb_typeof(validation_report) = 'object'),
  created_at timestamptz not null default now()
);
comment on table public.relevance_publications is 'Immutable version identity and provenance. Activation is stored separately.';

create or replace function public.valid_relevance_formula_config(p_config jsonb)
returns boolean language plpgsql immutable security invoker set search_path = '' as $$
declare
  v_minimum_score numeric;
  v_maximum_score numeric;
  v_minimum_factor numeric;
  v_maximum_factor numeric;
begin
  if p_config is null or jsonb_typeof(p_config) is distinct from 'object'
    or jsonb_typeof(p_config -> 'minimumScore') is distinct from 'number'
    or jsonb_typeof(p_config -> 'maximumScore') is distinct from 'number'
    or jsonb_typeof(p_config -> 'minimumFactor') is distinct from 'number'
    or jsonb_typeof(p_config -> 'maximumFactor') is distinct from 'number' then
    return false;
  end if;
  v_minimum_score := (p_config ->> 'minimumScore')::numeric;
  v_maximum_score := (p_config ->> 'maximumScore')::numeric;
  v_minimum_factor := (p_config ->> 'minimumFactor')::numeric;
  v_maximum_factor := (p_config ->> 'maximumFactor')::numeric;
  return v_minimum_score >= 1 and v_maximum_score <= 10 and v_minimum_score < v_maximum_score
    and v_minimum_factor >= 0.7 and v_maximum_factor <= 1.3 and v_minimum_factor <= v_maximum_factor;
exception when others then
  return false;
end;
$$;

create or replace function public.calculate_relevance_factor(p_score numeric, p_config jsonb)
returns numeric(5,4) language plpgsql immutable security invoker set search_path = '' as $$
declare
  v_minimum_score numeric;
  v_maximum_score numeric;
  v_minimum_factor numeric;
  v_maximum_factor numeric;
  v_clamped_score numeric;
begin
  if not public.valid_relevance_formula_config(p_config) then
    raise exception 'Invalid relevance formula config';
  end if;
  v_minimum_score := (p_config ->> 'minimumScore')::numeric;
  v_maximum_score := (p_config ->> 'maximumScore')::numeric;
  v_minimum_factor := (p_config ->> 'minimumFactor')::numeric;
  v_maximum_factor := (p_config ->> 'maximumFactor')::numeric;
  if p_score is null then raise exception 'Relevance score is required'; end if;
  v_clamped_score := least(v_maximum_score, greatest(v_minimum_score, p_score));
  return round(v_maximum_factor - (v_maximum_factor - v_minimum_factor)
    * (v_clamped_score - v_minimum_score) / (v_maximum_score - v_minimum_score), 4);
end;
$$;

create or replace function public.calculate_relevance_shadow_interval(
  p_previous_interval integer,
  p_question_count integer,
  p_correct_count integer,
  p_perceived_difficulty text,
  p_is_first_contact boolean,
  p_score numeric,
  p_config jsonb
)
returns integer language plpgsql immutable security invoker set search_path = '' as $$
declare
  v_minimum_score numeric;
  v_maximum_score numeric;
  v_minimum_factor numeric;
  v_maximum_factor numeric;
  v_clamped_score numeric;
  v_relevance_factor numeric;
  v_accuracy integer;
  v_engine_factor numeric;
  v_difficulty_factor numeric;
  v_raw_interval numeric;
begin
  if p_previous_interval is null or p_previous_interval <= 0
    or p_question_count is null or p_question_count <= 0
    or p_correct_count is null or p_correct_count < 0 or p_correct_count > p_question_count
    or p_perceived_difficulty is null or p_perceived_difficulty not in (
      'Muito fácil','Fácil','Médio','Difícil','Muito difícil'
    )
    or p_is_first_contact is null or p_score is null
    or not public.valid_relevance_formula_config(p_config) then
    raise exception 'Invalid relevance shadow interval inputs';
  end if;

  v_minimum_score := (p_config ->> 'minimumScore')::numeric;
  v_maximum_score := (p_config ->> 'maximumScore')::numeric;
  v_minimum_factor := (p_config ->> 'minimumFactor')::numeric;
  v_maximum_factor := (p_config ->> 'maximumFactor')::numeric;
  v_clamped_score := least(v_maximum_score, greatest(v_minimum_score, p_score));
  -- Keep full precision here, like the TypeScript candidate engine. The
  -- persisted factor is rounded separately to numeric(5,4) for its snapshot.
  v_relevance_factor := v_maximum_factor - (v_maximum_factor - v_minimum_factor)
    * (v_clamped_score - v_minimum_score) / (v_maximum_score - v_minimum_score);
  v_accuracy := round(greatest(0, p_correct_count)::numeric / p_question_count * 100)::integer;

  if p_is_first_contact then
    if p_question_count < 20 then
      v_raw_interval := case p_perceived_difficulty
        when 'Muito fácil' then 21
        when 'Fácil' then 18
        when 'Médio' then 14
        when 'Difícil' then 10
        when 'Muito difícil' then 7
      end;
    else
      v_raw_interval := case
        when v_accuracy >= 85 then 21
        when v_accuracy >= 70 then 14
        when v_accuracy >= 50 then 10
        else 7
      end;
    end if;
  elsif p_question_count < 20 then
    v_engine_factor := case p_perceived_difficulty
      when 'Muito fácil' then 6
      when 'Fácil' then 4.5
      when 'Médio' then 3
      when 'Difícil' then 1.3
      when 'Muito difícil' then 0.8
    end;
    v_raw_interval := p_previous_interval * v_engine_factor * v_relevance_factor;
  else
    v_engine_factor := case
      when v_accuracy >= 85 then 6
      when v_accuracy >= 70 then 3
      when v_accuracy >= 50 then 1.3
      else 0.8
    end;
    v_difficulty_factor := case p_perceived_difficulty
      when 'Muito fácil' then 1.2
      when 'Fácil' then 1.1
      when 'Médio' then 1
      when 'Difícil' then 0.9
      when 'Muito difícil' then 0.8
    end;
    v_raw_interval := p_previous_interval * v_engine_factor
      * v_relevance_factor * v_difficulty_factor;
  end if;

  return greatest(7, least(180, round(v_raw_interval)::integer));
end;
$$;

create table public.relevance_engine_activations (
  exam_code text primary key,
  publication_id uuid not null references public.relevance_publications(id) on delete restrict,
  mode text not null check (mode = 'shadow'),
  formula_version text not null check (formula_version = 'global-linear-bounded-v1'),
  formula_config jsonb not null check (public.valid_relevance_formula_config(formula_config)),
  activated_at timestamptz not null default now(),
  unique (publication_id, exam_code)
);

create table public.canonical_topics (
  topic_id text primary key check (topic_id ~ '^(CM|CIR|PED|GO|MP)-[A-Z0-9]+-[0-9]{{3}}$'),
  first_publication_id uuid not null references public.relevance_publications(id) on delete restrict,
  created_at timestamptz not null default now()
);

create table public.canonical_topic_versions (
  publication_id uuid not null references public.relevance_publications(id) on delete cascade,
  topic_id text not null references public.canonical_topics(topic_id) on delete restrict,
  canonical_name text not null,
  source_major_area text not null,
  major_area text not null check (major_area in ('Clínica Médica','Cirurgia','Ginecologia e Obstetrícia','Pediatria','Preventiva')),
  specialty text not null,
  editorial_signal text,
  confidence text not null check (confidence in ('alta','média')),
  rationale text,
  reference_year smallint not null check (reference_year between 2000 and 2200),
  primary key (publication_id, topic_id)
);

create table public.canonical_topic_aliases (
  publication_id uuid not null,
  topic_id text not null,
  alias text not null,
  normalized_alias text not null,
  primary key (publication_id, topic_id, normalized_alias),
  foreign key (publication_id, topic_id)
    references public.canonical_topic_versions(publication_id, topic_id) on delete cascade
);

create table public.topic_relevance_scores (
  publication_id uuid not null,
  topic_id text not null,
  exam_code text not null,
  score numeric(3,1) not null check (score between 1 and 10),
  reference_year smallint not null check (reference_year between 2000 and 2200),
  primary key (publication_id, topic_id, exam_code, reference_year),
  foreign key (publication_id, topic_id)
    references public.canonical_topic_versions(publication_id, topic_id) on delete cascade
);

create table public.course_catalog_releases (
  id uuid primary key,
  publication_id uuid not null references public.relevance_publications(id) on delete restrict,
  provider_code text not null,
  provider_name text not null,
  release_name text not null,
  source_document text not null,
  source_edition text not null,
  source_volume text not null,
  source_extracted_on date,
  source_sha256 text not null check (source_sha256 ~ '^[0-9a-f]{{64}}$'),
  status text not null check (status in ('draft','validated','active','retired')),
  created_at timestamptz not null default now(),
  unique (publication_id, provider_code),
  unique (id, publication_id)
);
create unique index course_catalog_one_active_release_per_provider_idx
  on public.course_catalog_releases(provider_code) where status = 'active';

create table public.course_catalog_item_tombstones (
  provider_code text not null,
  external_id text not null,
  publication_id uuid not null references public.relevance_publications(id) on delete restrict,
  reason text not null,
  created_at timestamptz not null default now(),
  primary key (provider_code, external_id)
);

create table public.course_catalog_items (
  id uuid primary key,
  release_id uuid not null,
  publication_id uuid not null,
  external_id text not null,
  catalog_area text not null,
  title text not null,
  source_order integer not null check (source_order > 0),
  major_area text not null check (major_area in ('Clínica Médica','Cirurgia','Ginecologia e Obstetrícia','Pediatria','Preventiva')),
  major_area_derivation text not null check (major_area_derivation = 'highest_weight_topic_experimental'),
  specialty text not null,
  specialty_derivation text not null check (specialty_derivation in ('unanimous_canonical_specialty','integrated_content')),
  effective_relevance numeric(3,1) not null check (effective_relevance between 1 and 10),
  workload_weight smallint not null check (workload_weight between 1 and 20),
  created_at timestamptz not null default now(),
  unique (release_id, external_id),
  unique (id, publication_id),
  foreign key (release_id, publication_id)
    references public.course_catalog_releases(id, publication_id) on delete cascade
);
comment on column public.course_catalog_items.major_area is 'Experimental UI label from the highest-weight canonical link; aggregate reports must use linked canonical topics.';

create table public.course_catalog_item_topics (
  item_id uuid not null,
  publication_id uuid not null,
  topic_id text not null,
  weight numeric(10,9) not null check (weight > 0 and weight <= 1),
  primary key (item_id, topic_id),
  foreign key (item_id, publication_id)
    references public.course_catalog_items(id, publication_id) on delete cascade,
  foreign key (publication_id, topic_id)
    references public.canonical_topic_versions(publication_id, topic_id) on delete restrict
);

create index canonical_topic_versions_search_idx on public.canonical_topic_versions(publication_id, major_area, specialty);
create index canonical_topics_first_publication_idx on public.canonical_topics(first_publication_id);
create index canonical_topic_versions_topic_idx on public.canonical_topic_versions(topic_id);
create index canonical_topic_aliases_search_idx on public.canonical_topic_aliases(publication_id, normalized_alias);
create index topic_relevance_scores_exam_idx on public.topic_relevance_scores(exam_code, reference_year, publication_id);
create index course_catalog_item_tombstones_publication_idx on public.course_catalog_item_tombstones(publication_id);
create index course_catalog_items_publication_idx on public.course_catalog_items(publication_id);
create index course_catalog_items_release_order_idx on public.course_catalog_items(release_id, source_order);
create index course_catalog_item_topics_topic_idx on public.course_catalog_item_topics(publication_id, topic_id);

create or replace function public.reject_tombstoned_catalog_item()
returns trigger language plpgsql security invoker set search_path = '' as $$
begin
  if exists (
    select 1
    from public.course_catalog_releases r
    join public.course_catalog_item_tombstones t
      on t.provider_code = r.provider_code and t.external_id = new.external_id
    where r.id = new.release_id
  ) then
    raise exception 'Catalog item identifier is tombstoned: %', new.external_id;
  end if;
  return new;
end;
$$;
create trigger reject_tombstoned_catalog_item
  before insert or update of release_id, external_id on public.course_catalog_items
  for each row execute function public.reject_tombstoned_catalog_item();
"""

    validation_report = {
        "topics": len(dataset.topics),
        "aliases": dataset.alias_count,
        "items": {provider: len(items) for provider, items in dataset.items.items()},
        "links": {provider: sum(len(item.links) for item in items) for provider, items in dataset.items.items()},
        "formulas_with_valid_cache": dataset.formula_count,
        "tombstones": ["medcof:MC-186"],
        "model_scope": {
            "exam_code": "GLOBAL",
            "item_and_snapshot_exam_code_dimension": False,
            "future_exam_specific_scores_require_model_evolution": True,
        },
    }
    publication_rows = [[
        publication_id, PUBLICATION_KEY, 2026, "validated",
        "0.7 * weighted_mean(topic_score) + 0.3 * max(topic_score)",
        "half_up_1_decimal", workbook_path.name, dataset.workbook_sha256,
        document_path.name, dataset.document_sha256, "1.2.0",
        json.dumps(validation_report, ensure_ascii=False, separators=(",", ":")),
    ]]
    topic_identity_rows = [[topic.topic_id, publication_id] for topic in dataset.topics.values()]
    topic_version_rows = [[
        publication_id, topic.topic_id, topic.canonical_name, topic.source_major_area,
        topic.major_area, topic.specialty, topic.editorial_signal, topic.confidence,
        topic.rationale, topic.reference_year,
    ] for topic in dataset.topics.values()]
    alias_rows = [[publication_id, topic.topic_id, alias, normalize_alias(alias)]
                  for topic in dataset.topics.values() for alias in topic.aliases]
    relevance_rows = [[publication_id, topic.topic_id, "GLOBAL", Decimal(topic.score), topic.reference_year]
                      for topic in dataset.topics.values()]
    release_rows = []
    for provider, metadata in RELEASE_META.items():
        release_rows.append([
            releases[provider], publication_id, metadata["provider_code"], provider,
            metadata["release_name"], metadata["source_document"], metadata["source_edition"],
            metadata["source_volume"], metadata["source_extracted_on"], dataset.workbook_sha256,
            "validated",
        ])
    tombstone_rows = [["medcof", "MC-186", publication_id, "Duplicate removed in v3; identifier must remain unused"]]
    item_rows: list[list[Any]] = []
    link_rows: list[list[Any]] = []
    for provider, provider_items in dataset.items.items():
        release_id = releases[provider]
        for item in provider_items:
            item_id = uuid_for("item", str(release_id), item.external_id)
            specialties = {dataset.topics[link.topic_id].specialty for link in item.links}
            item_rows.append([
                item_id, release_id, publication_id, item.external_id, item.catalog_area,
                item.title, item.source_order, item.major_area,
                "highest_weight_topic_experimental", item.specialty,
                "unanimous_canonical_specialty" if len(specialties) == 1 else "integrated_content",
                item.effective_score, len(item.links),
            ])
            for link in item.links:
                link_rows.append([item_id, publication_id, link.topic_id, link.weight])

    data_sql = "\n\n".join([
        insert_rows("public.relevance_publications", [
            "id", "version_key", "reference_year", "validation_state", "score_formula",
            "rounding_mode", "source_filename", "source_sha256",
            "technical_document_filename", "technical_document_sha256", "importer_version",
            "validation_report",
        ], publication_rows),
        insert_rows("public.canonical_topics", ["topic_id", "first_publication_id"], topic_identity_rows),
        insert_rows("public.canonical_topic_versions", [
            "publication_id", "topic_id", "canonical_name", "source_major_area", "major_area",
            "specialty", "editorial_signal", "confidence", "rationale", "reference_year",
        ], topic_version_rows),
        insert_rows("public.canonical_topic_aliases", ["publication_id", "topic_id", "alias", "normalized_alias"], alias_rows),
        insert_rows("public.topic_relevance_scores", ["publication_id", "topic_id", "exam_code", "score", "reference_year"], relevance_rows),
        insert_rows("public.course_catalog_releases", [
            "id", "publication_id", "provider_code", "provider_name", "release_name",
            "source_document", "source_edition", "source_volume", "source_extracted_on",
            "source_sha256", "status",
        ], release_rows),
        insert_rows("public.course_catalog_item_tombstones", ["provider_code", "external_id", "publication_id", "reason"], tombstone_rows),
        insert_rows("public.course_catalog_items", [
            "id", "release_id", "publication_id", "external_id", "catalog_area", "title",
            "source_order", "major_area", "major_area_derivation", "specialty",
            "specialty_derivation", "effective_relevance", "workload_weight",
        ], item_rows),
        insert_rows("public.course_catalog_item_topics", ["item_id", "publication_id", "topic_id", "weight"], link_rows),
    ])

    assertion_sql = """
-- Defense in depth: the importer already validates these invariants, and the
-- migration refuses the seed if the emitted SQL no longer preserves them.
do $$
begin
  if (select count(*) from public.canonical_topic_versions) <> 492 then
    raise exception 'Expected 492 canonical topic versions';
  end if;
  if (select count(*) from public.course_catalog_items) <> 920 then
    raise exception 'Expected 920 catalog items';
  end if;
  if (select count(*) from public.course_catalog_item_topics) <> 1235 then
    raise exception 'Expected 1235 item-topic links';
  end if;
  if exists (
    select 1 from public.course_catalog_item_topics group by item_id
    having abs(sum(weight) - 1::numeric) > 0.000001
  ) then
    raise exception 'Catalog item weights must sum to 1';
  end if;
  if exists (
    with calculated as (
      select i.id,
        round((0.7 * sum(l.weight * s.score) + 0.3 * max(s.score))::numeric, 1) as expected_score
      from public.course_catalog_items i
      join public.course_catalog_item_topics l on l.item_id = i.id
      join public.topic_relevance_scores s
        on s.publication_id = l.publication_id and s.topic_id = l.topic_id
        and s.exam_code = 'GLOBAL' and s.reference_year = 2026
      group by i.id
    )
    select 1 from calculated c join public.course_catalog_items i on i.id = c.id
    where i.effective_relevance <> c.expected_score
  ) then
    raise exception 'Effective relevance does not match the half-up formula';
  end if;
  if exists (
    select 1 from public.course_catalog_items i
    join public.course_catalog_releases r on r.id = i.release_id
    where r.provider_code = 'medcof' and i.external_id = 'MC-186'
  ) then
    raise exception 'MC-186 must remain tombstoned';
  end if;
  if exists (select 1 from public.course_catalog_releases where status = 'active')
    or exists (select 1 from public.relevance_engine_activations) then
    raise exception 'The v3 seed must not activate catalog or numeric relevance';
  end if;
end;
$$;
"""

    integration_sql = """
create schema if not exists relevance_internal;
revoke all on schema relevance_internal from public, anon, authenticated;

create or replace function relevance_internal.protect_validated_relevance_data()
returns trigger language plpgsql security definer set search_path = '' as $$
declare
  v_publication_id uuid;
  v_old_publication_id uuid;
  v_new_publication_id uuid;
begin
  if tg_op <> 'INSERT' then
    v_old_publication_id := coalesce(
      nullif(to_jsonb(old) ->> 'publication_id','')::uuid,
      nullif(to_jsonb(old) ->> 'first_publication_id','')::uuid
    );
  end if;
  if tg_op <> 'DELETE' then
    v_new_publication_id := coalesce(
      nullif(to_jsonb(new) ->> 'publication_id','')::uuid,
      nullif(to_jsonb(new) ->> 'first_publication_id','')::uuid
    );
  end if;
  if tg_op = 'UPDATE' and v_new_publication_id is distinct from v_old_publication_id then
    raise exception 'Relevance publication identity is immutable: %', tg_table_name;
  end if;
  v_publication_id := coalesce(v_new_publication_id, v_old_publication_id);
  if exists (
    select 1 from public.relevance_publications p
    where p.id in (v_old_publication_id, v_new_publication_id, v_publication_id)
      and p.validation_state = 'validated'
  ) then
    raise exception 'Validated relevance data is immutable: %', tg_table_name;
  end if;
  if tg_op = 'DELETE' then return old; end if;
  return new;
end;
$$;

create or replace function relevance_internal.protect_relevance_publication()
returns trigger language plpgsql security definer set search_path = '' as $$
begin
  if old.validation_state = 'validated' then
    raise exception 'Validated relevance publication is immutable';
  end if;
  if tg_op = 'DELETE' then return old; end if;
  return new;
end;
$$;

create or replace function relevance_internal.protect_course_catalog_release()
returns trigger language plpgsql security definer set search_path = '' as $$
declare
  v_publication_state text;
begin
  if tg_op = 'INSERT' then
    select p.validation_state into v_publication_state
    from public.relevance_publications p where p.id = new.publication_id;
    if v_publication_state = 'validated' then
      raise exception 'Validated catalog releases cannot be appended';
    end if;
    if v_publication_state is distinct from 'draft' or new.status <> 'draft' then
      raise exception 'New catalog releases require a draft publication and draft status';
    end if;
    return new;
  end if;
  if tg_op = 'DELETE' then
    if exists (select 1 from public.relevance_publications p where p.id = old.publication_id and p.validation_state = 'validated') then
      raise exception 'Validated catalog releases cannot be deleted';
    end if;
    return old;
  end if;
  if new.id is distinct from old.id or new.publication_id is distinct from old.publication_id
    or new.provider_code is distinct from old.provider_code or new.provider_name is distinct from old.provider_name
    or new.release_name is distinct from old.release_name or new.source_document is distinct from old.source_document
    or new.source_edition is distinct from old.source_edition or new.source_volume is distinct from old.source_volume
    or new.source_extracted_on is distinct from old.source_extracted_on
    or new.source_sha256 is distinct from old.source_sha256 or new.created_at is distinct from old.created_at then
    raise exception 'Catalog release metadata is immutable';
  end if;
  select p.validation_state into v_publication_state
  from public.relevance_publications p where p.id = old.publication_id;
  if new.status is distinct from old.status and (
    v_publication_state is distinct from 'validated' or not (
    (old.status = 'draft' and new.status = 'validated')
    or (old.status = 'validated' and new.status in ('active','retired'))
    or (old.status = 'active' and new.status in ('validated','retired'))
  )) then
    raise exception 'Invalid catalog release status transition: % -> %', old.status, new.status;
  end if;
  return new;
end;
$$;

create or replace function relevance_internal.validate_relevance_activation()
returns trigger language plpgsql security definer set search_path = '' as $$
begin
  if not exists (
    select 1 from public.relevance_publications p
    where p.id = new.publication_id and p.validation_state = 'validated'
  ) then
    raise exception 'Relevance activation requires a validated publication';
  end if;
  if not exists (
    select 1 from public.topic_relevance_scores s
    where s.publication_id = new.publication_id and s.exam_code = new.exam_code
  ) then
    raise exception 'Relevance activation has no matching exam scores';
  end if;
  return new;
end;
$$;

create trigger protect_relevance_publication
  before update or delete on public.relevance_publications
  for each row execute function relevance_internal.protect_relevance_publication();
create trigger protect_canonical_topics
  before insert or update or delete on public.canonical_topics
  for each row execute function relevance_internal.protect_validated_relevance_data();
create trigger protect_canonical_topic_versions
  before insert or update or delete on public.canonical_topic_versions
  for each row execute function relevance_internal.protect_validated_relevance_data();
create trigger protect_canonical_topic_aliases
  before insert or update or delete on public.canonical_topic_aliases
  for each row execute function relevance_internal.protect_validated_relevance_data();
create trigger protect_topic_relevance_scores
  before insert or update or delete on public.topic_relevance_scores
  for each row execute function relevance_internal.protect_validated_relevance_data();
create trigger protect_course_catalog_releases
  before insert or update or delete on public.course_catalog_releases
  for each row execute function relevance_internal.protect_course_catalog_release();
create trigger protect_course_catalog_item_tombstones
  before insert or update or delete on public.course_catalog_item_tombstones
  for each row execute function relevance_internal.protect_validated_relevance_data();
create trigger protect_course_catalog_items
  before insert or update or delete on public.course_catalog_items
  for each row execute function relevance_internal.protect_validated_relevance_data();
create trigger protect_course_catalog_item_topics
  before insert or update or delete on public.course_catalog_item_topics
  for each row execute function relevance_internal.protect_validated_relevance_data();
create trigger validate_relevance_activation
  before insert or update on public.relevance_engine_activations
  for each row execute function relevance_internal.validate_relevance_activation();

alter table public.question_blocks
  add column catalog_item_id uuid references public.course_catalog_items(id) on delete restrict,
  add column relevance_version_id uuid references public.relevance_publications(id) on delete restrict,
  add column relevance_score numeric(3,1) check (relevance_score is null or relevance_score between 1 and 10),
  add column relevance_factor numeric(5,4) check (relevance_factor is null or relevance_factor between 0.7 and 1.3),
  add column relevance_formula_version text check (relevance_formula_version is null or relevance_formula_version = 'global-linear-bounded-v1'),
  add column relevance_formula_config jsonb check (relevance_formula_config is null or jsonb_typeof(relevance_formula_config) = 'object'),
  add column relevance_shadow_interval_days integer check (relevance_shadow_interval_days is null or relevance_shadow_interval_days between 7 and 180),
  add constraint question_blocks_relevance_version_required_check check (
    (relevance_score is null and relevance_factor is null and relevance_formula_version is null
      and relevance_formula_config is null and relevance_shadow_interval_days is null)
    or relevance_version_id is not null
  ),
  add constraint question_blocks_relevance_calculation_complete_check check (
    (relevance_factor is null and relevance_shadow_interval_days is null)
    or (relevance_factor is not null and relevance_score is not null
      and relevance_formula_version is not null and relevance_formula_config is not null)
  ),
  add constraint question_blocks_catalog_source_all_or_none_check check (
    (catalog_item_id is null and relevance_version_id is null and relevance_score is null)
    or (catalog_item_id is not null and relevance_version_id is not null and relevance_score is not null)
  ),
  add constraint question_blocks_catalog_version_fk foreign key (catalog_item_id, relevance_version_id)
    references public.course_catalog_items(id, publication_id) on delete restrict;

alter table public.block_reviews
  add column catalog_item_id uuid references public.course_catalog_items(id) on delete restrict,
  add column relevance_version_id uuid references public.relevance_publications(id) on delete restrict,
  add column relevance_score numeric(3,1) check (relevance_score is null or relevance_score between 1 and 10),
  add column relevance_factor numeric(5,4) check (relevance_factor is null or relevance_factor between 0.7 and 1.3),
  add column relevance_formula_version text check (relevance_formula_version is null or relevance_formula_version = 'global-linear-bounded-v1'),
  add column relevance_formula_config jsonb check (relevance_formula_config is null or jsonb_typeof(relevance_formula_config) = 'object'),
  add column relevance_shadow_interval_days integer check (relevance_shadow_interval_days is null or relevance_shadow_interval_days between 7 and 180),
  add constraint block_reviews_relevance_version_required_check check (
    (relevance_score is null and relevance_factor is null and relevance_formula_version is null
      and relevance_formula_config is null and relevance_shadow_interval_days is null)
    or relevance_version_id is not null
  ),
  add constraint block_reviews_relevance_calculation_complete_check check (
    (relevance_factor is null and relevance_shadow_interval_days is null)
    or (relevance_factor is not null and relevance_score is not null
      and relevance_formula_version is not null and relevance_formula_config is not null)
  ),
  add constraint block_reviews_catalog_source_all_or_none_check check (
    (catalog_item_id is null and relevance_version_id is null and relevance_score is null)
    or (catalog_item_id is not null and relevance_version_id is not null and relevance_score is not null)
  ),
  add constraint block_reviews_catalog_version_fk foreign key (catalog_item_id, relevance_version_id)
    references public.course_catalog_items(id, publication_id) on delete restrict;

create index question_blocks_catalog_item_idx on public.question_blocks(user_id, catalog_item_id) where catalog_item_id is not null;
create index question_blocks_catalog_version_fk_idx on public.question_blocks(catalog_item_id, relevance_version_id) where catalog_item_id is not null;
create index block_reviews_relevance_version_idx on public.block_reviews(user_id, relevance_version_id, review_date desc) where relevance_version_id is not null;
create index block_reviews_catalog_version_fk_idx on public.block_reviews(catalog_item_id, relevance_version_id) where catalog_item_id is not null;

create or replace function public.validate_question_block_catalog_snapshot()
returns trigger language plpgsql security invoker set search_path = '' as $$
declare
  v_publication_id uuid;
  v_score numeric(3,1);
  v_review public.block_reviews;
  v_any_calculation boolean;
  v_calculation_changed boolean;
begin
  v_any_calculation := new.relevance_factor is not null or new.relevance_formula_version is not null
    or new.relevance_formula_config is not null or new.relevance_shadow_interval_days is not null;
  if tg_op = 'UPDATE' and old.catalog_item_id is not null and (
    new.catalog_item_id is distinct from old.catalog_item_id
    or new.relevance_version_id is distinct from old.relevance_version_id
    or new.relevance_score is distinct from old.relevance_score
  ) then
    raise exception 'Catalog identity and source relevance are immutable once attached';
  end if;
  -- An unchanged attached item is already protected by its FK and stored
  -- snapshot. Do not re-read a catalog that may now be retired and hidden by RLS.
  if new.catalog_item_id is not null and not (
    tg_op = 'UPDATE'
    and new.catalog_item_id is not distinct from old.catalog_item_id
    and new.relevance_version_id is not distinct from old.relevance_version_id
    and new.relevance_score is not distinct from old.relevance_score
  ) then
    select publication_id, effective_relevance into v_publication_id, v_score
    from public.course_catalog_items where id = new.catalog_item_id;
    if not found then raise exception 'Catalog item not found'; end if;
    if new.relevance_version_id is distinct from v_publication_id
      or new.relevance_score is distinct from v_score then
      raise exception 'Catalog item, relevance version and source score do not match';
    end if;
  end if;

  -- A block stores the latest completed review's candidate snapshot. New blocks
  -- may attach immutable source metadata, but cannot invent calculation output.
  if tg_op = 'INSERT' then
    if v_any_calculation then
      raise exception 'Block shadow calculation must come from a completed review';
    end if;
    return new;
  end if;

  v_calculation_changed := new.relevance_factor is distinct from old.relevance_factor
    or new.relevance_formula_version is distinct from old.relevance_formula_version
    or new.relevance_formula_config is distinct from old.relevance_formula_config
    or new.relevance_shadow_interval_days is distinct from old.relevance_shadow_interval_days;
  if v_calculation_changed then
    select * into v_review
    from public.block_reviews r
    where r.block_id = new.id and r.user_id = new.user_id
    order by r.created_at desc, r.id desc
    limit 1;
    if not found then raise exception 'Block shadow calculation has no completed review'; end if;
    if new.catalog_item_id is distinct from v_review.catalog_item_id
      or new.relevance_version_id is distinct from v_review.relevance_version_id
      or new.relevance_score is distinct from v_review.relevance_score
      or new.relevance_factor is distinct from v_review.relevance_factor
      or new.relevance_formula_version is distinct from v_review.relevance_formula_version
      or new.relevance_formula_config is distinct from v_review.relevance_formula_config
      or new.relevance_shadow_interval_days is distinct from v_review.relevance_shadow_interval_days then
      raise exception 'Block relevance calculation does not match its latest completed review';
    end if;
  end if;
  return new;
end;
$$;
create trigger validate_question_block_catalog_snapshot
  before insert or update of catalog_item_id, relevance_version_id, relevance_score, relevance_factor,
    relevance_formula_version, relevance_formula_config, relevance_shadow_interval_days, interval_days on public.question_blocks
  for each row execute function public.validate_question_block_catalog_snapshot();

create or replace function public.validate_block_review_relevance_snapshot()
returns trigger language plpgsql security invoker set search_path = '' as $$
declare
  v_block public.question_blocks;
  v_formula_version text;
  v_formula_config jsonb;
  v_factor numeric(5,4);
  v_shadow_interval integer;
  v_any_snapshot boolean;
  v_any_calculation boolean;
begin
  v_any_snapshot := new.catalog_item_id is not null or new.relevance_version_id is not null
    or new.relevance_score is not null or new.relevance_factor is not null
    or new.relevance_formula_version is not null or new.relevance_formula_config is not null
    or new.relevance_shadow_interval_days is not null;
  v_any_calculation := new.relevance_factor is not null or new.relevance_formula_version is not null
    or new.relevance_formula_config is not null or new.relevance_shadow_interval_days is not null;
  select * into v_block from public.question_blocks
    where id = new.block_id and user_id = new.user_id;
  if not found then raise exception 'Review block not found'; end if;
  -- First contacts intentionally do not run the relevance experiment. Legacy
  -- all-null rows and source-only snapshots both remain compatible.
  if new.contact_type = 'first_contact' then
    if new.catalog_item_id is not distinct from v_block.catalog_item_id
      and new.relevance_version_id is not distinct from v_block.relevance_version_id
      and new.relevance_score is not distinct from v_block.relevance_score
      and not v_any_calculation then
      return new;
    end if;
    if not v_any_snapshot then return new; end if;
    raise exception 'First-contact relevance snapshot does not match its block or contains a calculation';
  end if;
  if v_block.relevance_version_id is null and v_block.relevance_score is null and v_block.catalog_item_id is null then
    if v_any_snapshot then raise exception 'Review relevance snapshot has no source block metadata'; end if;
    return new;
  end if;
  if new.catalog_item_id is distinct from v_block.catalog_item_id
    or new.relevance_version_id is distinct from v_block.relevance_version_id
    or new.relevance_score is distinct from v_block.relevance_score then
    raise exception 'Review source relevance does not match its block';
  end if;
  if new.previous_interval_days is distinct from v_block.interval_days
    or new.importance is distinct from v_block.importance then
    raise exception 'Review engine inputs do not match its block';
  end if;
  select a.formula_version, a.formula_config into v_formula_version, v_formula_config
  from public.relevance_engine_activations a
  where a.exam_code = 'GLOBAL' and a.publication_id = v_block.relevance_version_id and a.mode = 'shadow';
  if not found then
    if v_any_calculation then raise exception 'No GLOBAL shadow relevance activation for this review'; end if;
    return new;
  end if;
  if not v_any_calculation or new.relevance_factor is null or new.relevance_formula_version is null
    or new.relevance_formula_config is null or new.relevance_shadow_interval_days is null then
    raise exception 'Complete server-derived shadow snapshot is required';
  end if;
  v_factor := public.calculate_relevance_factor(v_block.relevance_score, v_formula_config);
  v_shadow_interval := public.calculate_relevance_shadow_interval(
    v_block.interval_days,
    new.question_count,
    new.correct_count,
    new.perceived_difficulty,
    false,
    v_block.relevance_score,
    v_formula_config
  );
  if new.relevance_formula_version is distinct from v_formula_version
    or new.relevance_formula_config is distinct from v_formula_config
    or new.relevance_factor is distinct from v_factor
    or new.relevance_shadow_interval_days is distinct from v_shadow_interval then
    raise exception 'Review relevance calculation does not match the active shadow configuration';
  end if;
  return new;
end;
$$;
create trigger validate_block_review_relevance_snapshot
  before insert on public.block_reviews
  for each row execute function public.validate_block_review_relevance_snapshot();

create or replace function public.protect_block_review_relevance_snapshot()
returns trigger language plpgsql security invoker set search_path = '' as $$
begin
  if new.catalog_item_id is distinct from old.catalog_item_id
    or new.relevance_version_id is distinct from old.relevance_version_id
    or new.relevance_score is distinct from old.relevance_score
    or new.relevance_factor is distinct from old.relevance_factor
    or new.relevance_formula_version is distinct from old.relevance_formula_version
    or new.relevance_formula_config is distinct from old.relevance_formula_config
    or new.relevance_shadow_interval_days is distinct from old.relevance_shadow_interval_days then
    raise exception 'Historical relevance snapshots are immutable';
  end if;
  return new;
end;
$$;
create trigger protect_block_review_relevance_snapshot
  before update of catalog_item_id, relevance_version_id, relevance_score, relevance_factor,
    relevance_formula_version, relevance_formula_config, relevance_shadow_interval_days on public.block_reviews
  for each row execute function public.protect_block_review_relevance_snapshot();

-- Keeps every concurrency, idempotency, chronology and occurrence check from
-- 20260918230725 while adding immutable relevance snapshots in shadow mode.
create or replace function public.complete_block_review(p_block_id uuid,p_operation_id uuid,p_review jsonb,p_changes jsonb)
returns jsonb language plpgsql security invoker set search_path = '' as $$
declare
  v_user uuid := auth.uid();
  v_block public.question_blocks;
  v_review public.block_reviews;
  v_expected integer;
  v_profile public.student_profiles;
  v_review_date date;
  v_next_date date;
  v_planned_date date;
  v_interval integer;
  v_today date;
  v_catalog_item_id uuid;
  v_relevance_version_id uuid;
  v_relevance_score numeric(3,1);
  v_relevance_factor numeric(5,4);
  v_relevance_formula_version text;
  v_relevance_formula_config jsonb;
  v_relevance_shadow_interval_days integer;
begin
  if v_user is null or p_operation_id is null then raise exception 'Authentication and operation ID required'; end if;
  select * into v_profile from public.student_profiles where user_id = v_user for update;
  if not found then raise exception 'Student profile not found'; end if;
  select * into v_block from public.question_blocks where id = p_block_id and user_id = v_user for update;
  if not found then raise exception 'Block not found'; end if;
  select * into v_review from public.block_reviews where user_id = v_user and operation_id = p_operation_id;
  if found then
    if v_review.block_id <> p_block_id then raise exception 'Operation ID belongs to another block'; end if;
    return jsonb_build_object('block',to_jsonb(v_block),'review',to_jsonb(v_review),'replayed',true);
  end if;
  v_expected := (p_review ->> 'expected_repetitions')::integer;
  if v_expected is null or v_expected <> v_block.repetitions then raise exception 'Review changed; reload the block before saving'; end if;
  if p_changes ->> 'next_review_date' is distinct from p_review ->> 'new_next_review_date' then raise exception 'Inconsistent next review date'; end if;
  v_review_date := (p_review ->> 'review_date')::date;
  v_next_date := (p_review ->> 'new_next_review_date')::date;
  v_planned_date := (p_changes ->> 'planned_review_date')::date;
  v_interval := (p_changes ->> 'interval_days')::integer;
  v_today := (now() at time zone v_profile.timezone)::date;
  if v_review_date is null or v_review_date > v_today or v_review_date < coalesce(v_block.last_review_date,v_block.study_date) then
    raise exception 'Review date must be between the last contact and today';
  end if;
  if v_next_date is null or v_interval is null or v_interval not between 7 and 180 or v_next_date <> v_review_date + v_interval then
    raise exception 'Invalid review interval';
  end if;
  if v_planned_date is not null and v_planned_date < v_today then raise exception 'Planned review date cannot be in the past'; end if;

  -- Catalog identity, publication and source score are immutable block metadata.
  -- The current GLOBAL shadow activation is the sole source of calculation
  -- version/config/factor; client values are assertions, never inputs.
  v_catalog_item_id := v_block.catalog_item_id;
  v_relevance_version_id := v_block.relevance_version_id;
  v_relevance_score := v_block.relevance_score;
  if nullif(p_review ->> 'catalog_item_id','') is not null
    and (p_review ->> 'catalog_item_id')::uuid is distinct from v_catalog_item_id then
    raise exception 'Review catalog identity does not match the block';
  end if;
  if nullif(p_review ->> 'relevance_version_id','') is not null
    and (p_review ->> 'relevance_version_id')::uuid is distinct from v_relevance_version_id then
    raise exception 'Review relevance version does not match the block';
  end if;
  if nullif(p_review ->> 'relevance_score','') is not null
    and (p_review ->> 'relevance_score')::numeric is distinct from v_relevance_score then
    raise exception 'Review relevance score does not match the block';
  end if;

  v_relevance_factor := null;
  v_relevance_formula_version := null;
  v_relevance_formula_config := null;
  v_relevance_shadow_interval_days := null;
  if v_relevance_version_id is not null and v_relevance_score is not null then
    select a.formula_version, a.formula_config
      into v_relevance_formula_version, v_relevance_formula_config
    from public.relevance_engine_activations a
    where a.exam_code = 'GLOBAL' and a.publication_id = v_relevance_version_id and a.mode = 'shadow';
    if found then
      v_relevance_factor := public.calculate_relevance_factor(v_relevance_score, v_relevance_formula_config);
      v_relevance_shadow_interval_days := public.calculate_relevance_shadow_interval(
        v_block.interval_days,
        (p_review ->> 'question_count')::integer,
        (p_review ->> 'correct_count')::integer,
        p_review ->> 'perceived_difficulty',
        false,
        v_relevance_score,
        v_relevance_formula_config
      );
    end if;
  end if;
  if nullif(p_review ->> 'relevance_formula_version','') is not null
    and p_review ->> 'relevance_formula_version' is distinct from v_relevance_formula_version then
    raise exception 'Review formula version diverges from the active shadow configuration';
  end if;
  if p_review -> 'relevance_formula_config' is not null
    and p_review -> 'relevance_formula_config' <> 'null'::jsonb
    and p_review -> 'relevance_formula_config' is distinct from v_relevance_formula_config then
    raise exception 'Review formula config diverges from the active shadow configuration';
  end if;
  if nullif(p_review ->> 'relevance_factor','') is not null
    and (v_relevance_factor is null
      or abs((p_review ->> 'relevance_factor')::numeric - v_relevance_factor) > 0.0001) then
    raise exception 'Review factor diverges from the active shadow configuration';
  end if;
  if nullif(p_review ->> 'relevance_shadow_interval_days','') is not null
    and (p_review ->> 'relevance_shadow_interval_days')::integer is distinct from v_relevance_shadow_interval_days then
    raise exception 'Review shadow interval diverges from the server calculation';
  end if;

  insert into public.block_reviews(block_id,user_id,operation_id,review_date,question_count,correct_count,accuracy_percentage,perceived_difficulty,time_spent_minutes,
    sm2_grade_calculated,previous_next_review_date,new_next_review_date,contact_type,engine_version,calculation_mode,performance_band,importance,previous_interval_days,new_interval_days,priority_score,
    catalog_item_id,relevance_version_id,relevance_score,relevance_factor,relevance_formula_version,relevance_formula_config,relevance_shadow_interval_days)
    values(p_block_id,v_user,p_operation_id,(p_review->>'review_date')::date,(p_review->>'question_count')::integer,(p_review->>'correct_count')::integer,
      (p_review->>'accuracy_percentage')::integer,p_review->>'perceived_difficulty',(p_review->>'time_spent_minutes')::integer,(p_review->>'sm2_grade_calculated')::integer,
      v_block.next_review_date,(p_review->>'new_next_review_date')::date,'review',coalesce(p_review->>'engine_version',v_block.engine_version),
      p_review->>'calculation_mode',p_review->>'performance_band',v_block.importance,v_block.interval_days,(p_changes->>'interval_days')::integer,(p_review->>'priority_score')::numeric,
      v_catalog_item_id,v_relevance_version_id,v_relevance_score,v_relevance_factor,v_relevance_formula_version,v_relevance_formula_config,v_relevance_shadow_interval_days)
    returning * into v_review;
  -- Ensure a first migrated completion also leaves an auditable completed occurrence.
  insert into public.weekly_plan_items(user_id,block_id,revision_number,due_week_start,original_due_week_start,week_starts_on,planned_review_date,planning_source)
    select v_user,p_block_id,v_block.repetitions,w,w,extract(dow from w)::integer,v_block.planned_review_date,v_block.planning_source
    from (select public.planning_week_start(v_user,case when v_block.planning_source='manual' then coalesce(v_block.planned_review_date,v_block.next_review_date) else v_block.next_review_date end,p.week_starts_on) w
      from public.student_profiles p where p.user_id = v_user) due on conflict do nothing;
  update public.weekly_plan_items set status = 'completed',completed_at = now(),completed_review_id = v_review.id
    where user_id = v_user and block_id = p_block_id and revision_number = v_block.repetitions;
  update public.question_blocks set question_count = v_review.question_count,correct_count = v_review.correct_count,accuracy_percentage = v_review.accuracy_percentage,
    perceived_difficulty = v_review.perceived_difficulty,time_spent_minutes = v_review.time_spent_minutes,repetitions = repetitions + 1,
    interval_days = (p_changes->>'interval_days')::integer,next_review_date = v_review.new_next_review_date,last_review_date = v_review.review_date,
    planned_review_date = (p_changes->>'planned_review_date')::date,planning_source = 'automatic',backlog_since = null,backlog_urgency = null,
    pre_exam_review_requested = false,performance_band = v_review.performance_band,calculation_mode = coalesce(v_review.calculation_mode,calculation_mode),engine_version = v_review.engine_version,
    relevance_factor = v_review.relevance_factor,relevance_formula_version = v_review.relevance_formula_version,
    relevance_formula_config = v_review.relevance_formula_config,relevance_shadow_interval_days = v_review.relevance_shadow_interval_days,
    calendar_sync_status = case when calendar_sync_enabled then 'pending' else 'disabled' end,calendar_last_error = null
    where id = p_block_id and user_id = v_user returning * into v_block;
  return jsonb_build_object('block',to_jsonb(v_block),'review',to_jsonb(v_review),'replayed',false);
end;
$$;
"""

    access_sql = """
alter table public.relevance_publications enable row level security;
alter table public.relevance_engine_activations enable row level security;
alter table public.canonical_topics enable row level security;
alter table public.canonical_topic_versions enable row level security;
alter table public.canonical_topic_aliases enable row level security;
alter table public.topic_relevance_scores enable row level security;
alter table public.course_catalog_releases enable row level security;
alter table public.course_catalog_item_tombstones enable row level security;
alter table public.course_catalog_items enable row level security;
alter table public.course_catalog_item_topics enable row level security;

revoke all on public.relevance_publications, public.relevance_engine_activations,
  public.canonical_topics, public.canonical_topic_versions, public.canonical_topic_aliases,
  public.topic_relevance_scores, public.course_catalog_releases,
  public.course_catalog_item_tombstones, public.course_catalog_items,
  public.course_catalog_item_topics from anon, authenticated;
grant select on public.relevance_publications, public.relevance_engine_activations,
  public.canonical_topics, public.canonical_topic_versions, public.canonical_topic_aliases,
  public.topic_relevance_scores, public.course_catalog_releases,
  public.course_catalog_items, public.course_catalog_item_topics to authenticated;

create policy "Authenticated users can read active catalog releases"
  on public.course_catalog_releases for select to authenticated
  using (status = 'active');
create policy "Authenticated users can read active catalog items"
  on public.course_catalog_items for select to authenticated
  using (exists (select 1 from public.course_catalog_releases r where r.id = course_catalog_items.release_id and r.status = 'active'));
create policy "Authenticated users can read active catalog links"
  on public.course_catalog_item_topics for select to authenticated
  using (exists (
    select 1 from public.course_catalog_items i
    join public.course_catalog_releases r on r.id = i.release_id
    where i.id = course_catalog_item_topics.item_id and r.status = 'active'
  ));
create policy "Authenticated users can read activated engine modes"
  on public.relevance_engine_activations for select to authenticated
  using (mode = 'shadow');
create policy "Authenticated users can read activated publications"
  on public.relevance_publications for select to authenticated
  using (
    exists (select 1 from public.course_catalog_releases r where r.publication_id = relevance_publications.id and r.status = 'active')
    or exists (select 1 from public.relevance_engine_activations a where a.publication_id = relevance_publications.id and a.mode = 'shadow')
  );
create policy "Authenticated users can read topics in active catalogs"
  on public.canonical_topics for select to authenticated
  using (exists (
    select 1 from public.course_catalog_item_topics l
    join public.course_catalog_items i on i.id = l.item_id
    join public.course_catalog_releases r on r.id = i.release_id
    where l.topic_id = canonical_topics.topic_id and r.status = 'active'
  ));
create policy "Authenticated users can read topic versions in active catalogs"
  on public.canonical_topic_versions for select to authenticated
  using (exists (
    select 1 from public.course_catalog_item_topics l
    join public.course_catalog_items i on i.id = l.item_id
    join public.course_catalog_releases r on r.id = i.release_id
    where l.topic_id = canonical_topic_versions.topic_id
      and l.publication_id = canonical_topic_versions.publication_id and r.status = 'active'
  ));
create policy "Authenticated users can read aliases in active catalogs"
  on public.canonical_topic_aliases for select to authenticated
  using (exists (
    select 1 from public.course_catalog_item_topics l
    join public.course_catalog_items i on i.id = l.item_id
    join public.course_catalog_releases r on r.id = i.release_id
    where l.topic_id = canonical_topic_aliases.topic_id
      and l.publication_id = canonical_topic_aliases.publication_id and r.status = 'active'
  ));
create policy "Authenticated users can read scores only in activated engine modes"
  on public.topic_relevance_scores for select to authenticated
  using (exists (
    select 1 from public.relevance_engine_activations a
    where a.publication_id = topic_relevance_scores.publication_id
      and a.exam_code = topic_relevance_scores.exam_code and a.mode = 'shadow'
  ));

create or replace function public.normalize_catalog_search(p_value text)
returns text language sql immutable security invoker set search_path = '' as $$
  select translate(
    lower(regexp_replace(trim(coalesce(p_value,'')), '[[:space:]]+', ' ', 'g')),
    'áàâãäéèêëíìîïóòôõöúùûüçñ',
    'aaaaaeeeeiiiiooooouuuucn'
  );
$$;

create or replace function public.get_catalog_suggestions(
  p_release_id uuid,
  p_search text default null,
  p_limit integer default 200,
  p_offset integer default 0
)
returns table (
  catalog_release_id uuid,
  catalog_item_id uuid,
  external_id text,
  title text,
  catalog_area text,
  major_area text,
  specialty text,
  effective_score numeric(3,1),
  relevance_version_id uuid,
  topic_ids text[],
  workload_weight integer,
  suggested_importance text
)
language plpgsql stable security invoker set search_path = '' as $$
begin
  if auth.uid() is null then raise exception 'Authentication required'; end if;
  if p_release_id is null then raise exception 'Catalog release is required'; end if;
  if p_limit is null or p_limit not between 1 and 500 then raise exception 'Limit must be between 1 and 500'; end if;
  if p_offset is null or p_offset not between 0 and 1000000 then raise exception 'Offset must be between 0 and 1000000'; end if;
  return query
    select i.release_id, i.id, i.external_id, i.title, i.catalog_area, i.major_area, i.specialty,
      i.effective_relevance, i.publication_id,
      array_agg(l.topic_id order by l.topic_id), i.workload_weight::integer,
      case when i.effective_relevance >= 7.5 then 'alta'
           when i.effective_relevance <= 3.5 then 'baixa' else 'media' end
    from public.course_catalog_items i
    join public.course_catalog_releases r on r.id = i.release_id and r.status = 'active'
    join public.course_catalog_item_topics l on l.item_id = i.id
    where i.release_id = p_release_id
      and (
        nullif(trim(p_search),'') is null
        or i.title ilike '%' || trim(p_search) || '%'
        or i.catalog_area ilike '%' || trim(p_search) || '%'
        or i.major_area ilike '%' || trim(p_search) || '%'
        or i.specialty ilike '%' || trim(p_search) || '%'
        or exists (
          select 1
          from public.course_catalog_item_topics search_link
          join public.canonical_topic_aliases a
            on a.publication_id = search_link.publication_id and a.topic_id = search_link.topic_id
          where search_link.item_id = i.id
            and a.normalized_alias like '%' || public.normalize_catalog_search(p_search) || '%'
        )
      )
    group by i.release_id, i.id
    order by i.source_order, i.external_id
    limit p_limit offset p_offset;
end;
$$;

revoke execute on function public.reject_tombstoned_catalog_item(),
  public.validate_question_block_catalog_snapshot(),
  public.validate_block_review_relevance_snapshot(),
  public.protect_block_review_relevance_snapshot(),
  public.get_catalog_suggestions(uuid,text,integer,integer) from public, anon;
revoke execute on function relevance_internal.protect_validated_relevance_data(),
  relevance_internal.protect_relevance_publication(),
  relevance_internal.protect_course_catalog_release(),
  relevance_internal.validate_relevance_activation() from public, anon, authenticated;
grant execute on function public.get_catalog_suggestions(uuid,text,integer,integer) to authenticated;

-- Deliberate activation boundary:
-- * Catalog: a reviewed follow-up migration may set exactly one provider release to active.
-- * Numeric motor: a reviewed follow-up migration may insert GLOBAL in shadow mode.
--   Active interval semantics are deliberately outside this foundation.
-- This migration performs neither action.
"""
    return schema + "\n" + data_sql + "\n\n" + assertion_sql + "\n" + integration_sql + "\n" + access_sql


def audit_summary(dataset: Dataset) -> dict[str, Any]:
    def factor(score: Decimal) -> Decimal:
        return (Decimal("1.2") - Decimal("0.4") * (score - Decimal("1")) / Decimal("9")).quantize(
            Decimal("0.0001"), rounding=ROUND_HALF_UP
        )

    topic_scores = Counter(str(topic.score) for topic in dataset.topics.values())
    item_scores = Counter(str(item.effective_score) for items in dataset.items.values() for item in items)
    item_factors = Counter(str(factor(item.effective_score)) for items in dataset.items.values() for item in items)
    return {
        "publication_key": PUBLICATION_KEY,
        "workbook_sha256": dataset.workbook_sha256,
        "technical_document_sha256": dataset.document_sha256,
        "topics": len(dataset.topics),
        "aliases": dataset.alias_count,
        "items": {provider: len(items) for provider, items in dataset.items.items()},
        "links": {provider: sum(len(item.links) for item in items) for provider, items in dataset.items.items()},
        "formulas_with_valid_cache": dataset.formula_count,
        "catalog_status": "validated",
        "numeric_engine_activation": None,
        "tombstones": ["medcof:MC-186"],
        "known_model_limitations": [
            "Item effective relevance and block/review snapshots are GLOBAL-only and do not carry exam_code; future exam-specific scores require a new publication and model evolution."
        ],
        "experimental_shadow_policy": {
            "config_id": "global-linear-bounded-v1",
            "formula": "1.2 - 0.4 * (R - 1) / 9",
            "status": "experimental_inactive",
            "topic_score_distribution": dict(sorted(topic_scores.items(), key=lambda pair: Decimal(pair[0]))),
            "item_effective_score_distribution": dict(sorted(item_scores.items(), key=lambda pair: Decimal(pair[0]))),
            "item_factor_distribution": dict(sorted(item_factors.items(), key=lambda pair: Decimal(pair[0]))),
        },
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--technical-document", type=Path, required=True)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--audit-json", type=Path)
    parser.add_argument("--check-only", action="store_true")
    args = parser.parse_args()
    if args.check_only and args.output:
        parser.error("--check-only cannot be combined with --output")
    try:
        dataset = load_dataset(args.workbook, args.technical_document)
        summary = audit_summary(dataset)
        if args.audit_json:
            args.audit_json.parent.mkdir(parents=True, exist_ok=True)
            args.audit_json.write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        if args.output:
            args.output.parent.mkdir(parents=True, exist_ok=True)
            args.output.write_text(render_migration(dataset, args.workbook, args.technical_document), encoding="utf-8", newline="\n")
        print(json.dumps(summary, ensure_ascii=False, indent=2))
        return 0
    except ValidationError as exc:
        print(f"VALIDATION FAILED: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
